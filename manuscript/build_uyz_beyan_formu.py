#!/usr/bin/env python3
"""KOÜ Sosyal Bilimler Enstitüsü ÜYZ Beyan Formu — Word üretici."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
LOGO_SRC = ASSETS / "kou_logo_0.png"
LOGO_WM = ASSETS / "kou_logo_watermark.png"
OUTPUT = ROOT / "KOÜ_SBE_ÜYZ_Kullanımı_Beyan_Formu.docx"

FONT = "Times New Roman"
SIZE = Pt(12)
BLACK = RGBColor(0, 0, 0)


def make_faded_logo(src: Path, dest: Path, opacity: float = 0.08) -> Path:
    img = Image.open(src).convert("RGBA")
    pixels = img.load()
    w, h = img.size
    for y in range(h):
        for x in range(w):
            r, g, b, a = pixels[x, y]
            pixels[x, y] = (r, g, b, int(a * opacity))
    dest.parent.mkdir(parents=True, exist_ok=True)
    img.save(dest, "PNG")
    return dest


def set_run_font(run, bold: bool = False, italic: bool = False, size=SIZE) -> None:
    run.font.name = FONT
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:cs"), FONT)
    rPr.insert(0, rFonts)


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    space_after: int = 6,
    space_before: int = 0,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        run = p.add_run(text)
        set_run_font(run, bold=bold, italic=italic)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    sizes = {1: Pt(14), 2: Pt(12), 3: Pt(12)}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, bold=True, size=sizes.get(level, SIZE))


def style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    set_run_font(run)
            if not cell.paragraphs[0].runs:
                set_run_font(cell.paragraphs[0].add_run())


def fill_label_row(table, row_idx: int, label: str, placeholder: str = "") -> None:
    row = table.rows[row_idx]
    row.cells[0].text = label
    row.cells[1].text = placeholder
    style_table(table)


def add_checkbox_line(doc: Document, text: str) -> None:
    add_paragraph(doc, f"☐ {text}", space_after=4)


def build_document() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    style.font.color.rgb = BLACK

    add_paragraph(
        doc,
        "KOCAELİ ÜNİVERSİTESİ",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=0,
    )
    add_paragraph(
        doc,
        "SOSYAL BİLİMLER ENSTİTÜSÜ",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_paragraph(
        doc,
        "ÜRETKEN YAPAY ZEKÂ (ÜYZ) KULLANIMI BEYAN FORMU",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )

    add_paragraph(
        doc,
        "Bu form, Enstitümüzde hazırlanan yüksek lisans ve doktora tezlerinde üretken yapay "
        "zekâ araçlarının nasıl kullanıldığının açıkça belirtilmesi için düzenlenmiştir. "
        "Öğrenci ve danışman tarafından imzalanır. Danışman, Turnitin yapay zekâ kullanım "
        "oranını da forma işler. Etik Beyan sayfasından hemen sonra eklenir; tez metninde "
        "Kaynaklar bölümünden önce kısa bir beyan metni yer alır.",
        space_after=10,
    )

    add_heading(doc, "A. Öğrenci ve tez bilgileri", 2)
    t1 = doc.add_table(rows=7, cols=2)
    t1.style = "Table Grid"
    labels = [
        ("Adı Soyadı", ""),
        ("Öğrenci No", ""),
        ("Anabilim Dalı / Program", ""),
        ("Tez Türü", "☐ Yüksek Lisans    ☐ Doktora"),
        ("Tez Başlığı", ""),
        ("Danışman", ""),
        ("Tez Teslim Tarihi", ""),
    ]
    for i, (lab, val) in enumerate(labels):
        fill_label_row(t1, i, lab, val)
    style_table(t1)
    add_paragraph(doc, "", space_after=6)

    add_heading(doc, "B. Genel beyan", 2)
    add_paragraph(doc, "Aşağıdaki maddeleri okuyup onayladıktan sonra imzalayınız.", space_after=6)
    statements = [
        "Bu tezde kullandığım tüm üretken yapay zekâ araçlarını bu formda ve tez metninde açıkça belirttim.",
        "Yapay zekâ araçlarını yazar veya ortak yazar olarak göstermedim.",
        "Yapay zekâ çıktılarının doğruluğunu, tarafsızlığını ve kaynakların gerçekliğini kendim kontrol ettim.",
        "Katılımcı verilerini, kişisel verileri ve gizli araştırma materyallerini yapay zekâ araçlarına yüklemedim.",
        "Tezin bilimsel katkısı, analizi ve yorumu bana aittir.",
        "Bu formda yer alan yasak kullanım hükümlerine aykırı bir işlem yapmadım.",
    ]
    for s in statements:
        add_checkbox_line(doc, s)
    add_paragraph(doc, "", space_after=6)

    add_heading(doc, "C. Kullanım durumu (birini işaretleyiniz)", 2)
    add_checkbox_line(
        doc,
        "Seçenek 1 — Bu tezde üretken yapay zekâ kullanılmamıştır. "
        "(D ve E bölümleri doldurulmaz.)",
    )
    add_checkbox_line(
        doc,
        "Seçenek 2 — Sınırlı düzeyde kullanım yapılmıştır. (Tablo 1 doldurulur.)",
    )
    add_checkbox_line(
        doc,
        "Seçenek 3 — Kapsamlı kullanım yapılmıştır. (Tablo 1 ve Tablo 2 doldurulur.)",
    )
    add_paragraph(doc, "", space_after=6)

    add_heading(doc, "D. Üretken yapay zekâ kullanım kaydı", 2)
    add_paragraph(doc, "Tablo 1 — Kullanılan araçlar", bold=True, space_after=4)
    t2 = doc.add_table(rows=4, cols=7)
    t2.style = "Table Grid"
    headers = [
        "Sıra",
        "Araç (ad ve sürüm)",
        "Kullanım tarihi",
        "Tez bölümü / aşaması",
        "Kullanım amacı",
        "Düzey\n(2 / 3)",
        "Doğrulama",
    ]
    for i, h in enumerate(headers):
        t2.rows[0].cells[i].text = h
    for r in range(1, 4):
        t2.rows[r].cells[0].text = str(r)
    style_table(t2)
    add_paragraph(
        doc,
        "Tez aşaması örnekleri: literatür taraması, yöntem, veri toplama, veri analizi, "
        "bulgular, tartışma, özet, görsel-şekil, dil düzenleme.",
        italic=True,
        space_after=8,
    )

    add_paragraph(doc, "Tablo 2 — Kapsamlı kullanımlar için (Seçenek 3)", bold=True, space_after=4)
    t3 = doc.add_table(rows=3, cols=5)
    t3.style = "Table Grid"
    h2 = [
        "Sıra",
        "Yapay zekânın ürettiği çıktı",
        "Yazarın katkısı ve düzenleme süreci",
        "Doğrulama yöntemi",
        "Tezde beyan yeri",
    ]
    for i, h in enumerate(h2):
        t3.rows[0].cells[i].text = h
    for r in range(1, 3):
        t3.rows[r].cells[0].text = str(r)
    style_table(t3)
    add_paragraph(doc, "", space_after=6)

    add_paragraph(doc, "Tablo 3 — Bölümlere göre özet (isteğe bağlı)", bold=True, space_after=4)
    t4 = doc.add_table(rows=9, cols=4)
    t4.style = "Table Grid"
    t4.rows[0].cells[0].text = "Tez bölümü"
    t4.rows[0].cells[1].text = "Kullanıldı mı?"
    t4.rows[0].cells[2].text = "Düzey"
    t4.rows[0].cells[3].text = "Kısa açıklama"
    bolumler = [
        "Özet / Abstract",
        "Giriş",
        "Kuramsal çerçeve / Literatür",
        "Yöntem",
        "Bulgular",
        "Tartışma",
        "Sonuç ve öneriler",
        "Ekler / Görseller",
    ]
    for i, b in enumerate(bolumler, start=1):
        t4.rows[i].cells[0].text = b
        t4.rows[i].cells[1].text = "☐ Hayır   ☐ Evet"
    style_table(t4)
    add_paragraph(doc, "", space_after=8)

    add_heading(doc, "E. Tez metnine eklenecek beyan metni", 2)
    add_paragraph(
        doc,
        "Aşağıdaki metin, tezde Kaynaklar bölümünden önce yer alır. Tablo 1–2’ye göre "
        "doldurulur.",
        space_after=6,
    )
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    cell = box.rows[0].cells[0]
    cell.text = (
        "Üretken Yapay Zekâ Kullanım Beyanı\n\n"
        "[Bu çalışmanın hazırlanması sürecinde … aracı … amacıyla kullanılmıştır. "
        "Üretilen çıktılar yazar tarafından gözden geçirilmiş, doğrulanmış ve "
        "yeniden düzenlenmiştir. Yapay zekâ araçları yazar olarak gösterilmemiştir. "
        "Nihai metinden doğan sorumluluk yazara aittir.]"
    )
    for p in cell.paragraphs:
        for run in p.runs:
            set_run_font(run, italic=True)
    style_table(box)
    add_paragraph(doc, "", space_after=8)

    add_heading(doc, "F. Öğrenci beyanı ve imza", 2)
    add_paragraph(
        doc,
        "Bu formda verdiğim bilgilerin doğru olduğunu; yapay zekâ kullanımına ilişkin "
        "tüm sorumluluğun bana ait olduğunu beyan ederim.",
        space_after=6,
    )
    t5 = doc.add_table(rows=2, cols=3)
    t5.style = "Table Grid"
    t5.rows[0].cells[0].text = "Öğrenci adı soyadı"
    t5.rows[0].cells[1].text = "İmza"
    t5.rows[0].cells[2].text = "Tarih"
    style_table(t5)
    add_paragraph(doc, "", space_after=8)

    add_heading(doc, "G. Danışman beyanı ve imza", 2)
    add_paragraph(
        doc,
        "Bu bölüm tez danışmanı tarafından doldurulur ve imzalanır.",
        space_after=6,
    )
    advisor_statements = [
        "Öğrencinin bu formda ve tez metninde yaptığı yapay zekâ kullanım beyanını inceledim.",
        "Tez metni Turnitin üzerinden yapay zekâ kullanım oranı açısından değerlendirilmiştir.",
        "Öğrencinin beyanı ile Turnitin sonucu arasında belirgin bir uyumsuzluk bulunmamaktadır.",
        "Tezdeki yapay zekâ kullanımının Enstitü kuralları ve bilimsel etik ilkeleriyle bağdaştığını onaylarım.",
    ]
    for s in advisor_statements:
        add_checkbox_line(doc, s)
    add_paragraph(doc, "", space_after=6)

    add_paragraph(doc, "Turnitin yapay zekâ kullanım raporu", bold=True, space_after=4)
    t_turnitin = doc.add_table(rows=4, cols=2)
    t_turnitin.style = "Table Grid"
    turnitin_rows = [
        ("Turnitin yapay zekâ kullanım oranı (%)", ""),
        ("Turnitin rapor tarihi", ""),
        ("Değerlendirilen dosya / sürüm", ""),
        ("Rapor referansı veya not (isteğe bağlı)", ""),
    ]
    for i, (lab, val) in enumerate(turnitin_rows):
        fill_label_row(t_turnitin, i, lab, val)
    style_table(t_turnitin)
    add_paragraph(
        doc,
        "Oran, danışman tarafından Turnitin sisteminden alınan güncel yapay zekâ "
        "kullanım raporuna göre yazılır.",
        italic=True,
        space_after=8,
    )

    t6 = doc.add_table(rows=2, cols=3)
    t6.style = "Table Grid"
    t6.rows[0].cells[0].text = "Danışman adı soyadı"
    t6.rows[0].cells[1].text = "İmza"
    t6.rows[0].cells[2].text = "Tarih"
    style_table(t6)

    doc.add_page_break()

    add_heading(doc, "EK — Kullanım düzeyleri ve kurallar", 1)
    add_paragraph(
        doc,
        "Aşağıdaki özet, formu doldururken başvurulacak ortak çerçevedir. Ayrıntılar "
        "Enstitü tez yazım kılavuzunda yer alır.",
        space_after=8,
    )

    add_heading(doc, "1. Beyan gerektirmeyen kullanımlar", 2)
    add_paragraph(doc, "• Yazım ve dilbilgisi denetimi (içerik değiştirilmeden)", space_after=2)
    add_paragraph(doc, "• Kaynakça biçimlendirme (Zotero, Mendeley vb.)", space_after=6)

    add_heading(doc, "2. Kısa beyan gerektiren kullanımlar", 2)
    items_b = [
        "Metnin dil ve üslubunun düzeltilmesi (anlam korunarak)",
        "Çeviri veya dil kontrolü",
        "Araştırma sorusu ve kavramsal çerçeve için fikir geliştirme",
        "Literatür taramasında yönlendirme (referanslar ayrıca doğrulanmalı)",
        "Veri analizi kodunda yardım (kod test edilmeli)",
    ]
    for it in items_b:
        add_paragraph(doc, f"• {it}", space_after=2)
    add_paragraph(doc, "", space_after=4)

    add_heading(doc, "3. Ayrıntılı beyan gerektiren kullanımlar", 2)
    items_c = [
        "Tez bölümü veya alt bölüm taslağı (kapsamlı yeniden yazım şart)",
        "Nitel veya nicel veri özetleme / kod önerisi (insan doğrulaması şart)",
        "Anket, görüşme rehberi veya ölçek maddesi taslağı (danışman onayı şart)",
        "Yapay zekâ ile görsel, tablo veya şekil üretimi",
    ]
    for it in items_c:
        add_paragraph(doc, f"• {it}", space_after=2)
    add_paragraph(doc, "", space_after=4)

    add_heading(doc, "4. Yasak kullanımlar", 2)
    items_d = [
        "Tezin tamamının veya temel bölümlerinin yapay zekâya yazdırılması",
        "Gerçek veri veya katılımcı uydurma; sentetik veriyi gerçek veri gibi sunma",
        "Var olmayan kaynak veya atıf kullanma",
        "Yapay zekâ çıktısını kendi çalışması gibi gösterme veya kullanımı gizleme",
        "Görüşme transkripti, kişisel veri ve gizli materyallerin yapay zekâ araçlarına verilmesi",
        "Jüri ve değerlendirme süreçlerinde yapay zekâ kullanımı",
    ]
    for it in items_d:
        add_paragraph(doc, f"• {it}", space_after=2)
    add_paragraph(doc, "", space_after=6)

    add_heading(doc, "5. Dayanak", 2)
    add_paragraph(
        doc,
        "YÖK, Yükseköğretim Kurumları Bilimsel Araştırma ve Yayın Faaliyetlerinde "
        "Üretken Yapay Zekâ Kullanımına Dair Etik Rehber (2024); Yükseköğretim Kurumları "
        "Bilimsel Araştırma ve Yayın Etiği Yönergesi; 6698 sayılı KVKK.",
        space_after=6,
    )

    add_paragraph(
        doc,
        f"Form sürümü: 1.0    |    Kocaeli Üniversitesi Sosyal Bilimler Enstitüsü",
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    return doc


def add_image_watermark(doc_path: Path, image_path: Path) -> None:
    """Word belgesine sayfa filigranı (logo) ekler."""
    from docx import Document as Doc

    doc = Doc(doc_path)
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False

    # Mevcut header içeriğini temizle
    for child in list(header._element):
        header._element.remove(child)

    paragraph = header.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(3.8))

    # Görseli filigran katmanına taşımak için VML şekli ekle
    hdr = header._element
    pict = parse_xml(
        f"""
        <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
             xmlns:v="urn:schemas-microsoft-com:vml"
             xmlns:o="urn:schemas-microsoft-com:office:office"
             xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
          <w:pPr><w:pStyle w:val="Header"/></w:pPr>
          <w:r>
            <w:pict>
              <v:shapetype id="_x0000_t75" coordsize="21600,21600"
                o:spt="75" o:preferrelative="t" path="m@4@3l@4@0@0@0@0@0@0" filled="f" stroked="f">
                <v:stroke joinstyle="miter"/>
                <v:formulas>
                  <v:f eqn="if lineDrawn pixelLineWidth 0"/>
                  <v:f eqn="sum @0 1 0"/>
                  <v:f eqn="sum 0 0 @1"/>
                  <v:f eqn="prod @2 1 2"/>
                  <v:f eqn="prod @3 21600 pixelWidth"/>
                  <v:f eqn="prod @3 21600 pixelHeight"/>
                  <v:f eqn="sum @0 0 1"/>
                  <v:f eqn="prod @6 1 2"/>
                  <v:f eqn="prod @7 21600 pixelWidth"/>
                  <v:f eqn="sum @8 21600 0"/>
                  <v:f eqn="prod @7 21600 pixelHeight"/>
                  <v:f eqn="sum @10 21600 0"/>
                </v:formulas>
                <v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/>
                <o:lock v:ext="edit" aspectratio="t"/>
              </v:shapetype>
              <v:shape id="Watermark" type="#_x0000_t75"
                style="position:absolute;margin-left:0;margin-top:0;width:420pt;height:420pt;z-index:-251658240;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
                o:allowincell="f">
                <v:imagedata r:id="rIdWatermark" o:title="KOÜ"/>
              </v:shape>
            </w:pict>
          </w:r>
        </w:p>
        """
    )
    # python-docx ile ilişki kur
    part = header.part
    with open(image_path, "rb") as f:
        image_stream = io.BytesIO(f.read())
    image_part = part.package.get_or_add_image_part(image_stream)
    r_id = part.relate_to(image_part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")

    for imagedata in pict.iter("{urn:schemas-microsoft-com:vml}imagedata"):
        imagedata.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", r_id)

    # Eski paragrafı kaldır, VML paragrafını ekle
    for child in list(hdr):
        hdr.remove(child)
    hdr.append(pict)

    doc.save(doc_path)


def main() -> None:
    if not LOGO_SRC.exists():
        raise FileNotFoundError(f"Logo bulunamadı: {LOGO_SRC}")

    make_faded_logo(LOGO_SRC, LOGO_WM, opacity=0.10)
    doc = build_document()
    doc.save(OUTPUT)
    add_image_watermark(OUTPUT, LOGO_WM)
    print(f"Oluşturuldu: {OUTPUT}")


if __name__ == "__main__":
    main()
