#!/usr/bin/env python3
"""Build IJAL submission folder on Desktop with title page and blind manuscript."""

from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

DESKTOP = Path.home() / "Desktop"
SOURCE = DESKTOP / "IJAL_v1.docx"
OUTPUT_DIR = DESKTOP / "IJAL"

TITLE = (
    "Perceived External Discouragement from Significant Others in the "
    "Avoidance of English-Medium Instruction"
)
KEYWORDS = (
    "English-medium instruction, interdependent self-construal, linguistic capital, "
    "perceived external discouragement, self-determination theory, the ought-to L2 self"
)

DATA_AVAILABILITY = (
    "The survey dataset and interview transcripts supporting the findings of this study "
    "are not publicly available because participants did not consent to open sharing of "
    "their responses. De-identified excerpts and aggregate descriptive statistics are "
    "reported in the manuscript. Further information is available from the corresponding "
    "author upon reasonable request, subject to institutional and ethical constraints."
)
FUNDING = (
    "This research received no specific grant from any funding agency in the public, "
    "commercial, or not-for-profit sectors."
)
CONFLICT_OF_INTEREST = "The author declares no conflicts of interest."
ETHICS_APPROVAL = (
    "Participation was voluntary and anonymous, and informed consent was obtained in "
    "writing. Because the study relied on non-invasive self-report, ethical approval "
    "followed institutional guidelines for educational survey research at the time of "
    "data collection at Kocaeli University."
)
PERMISSION_TO_REPRODUCE = (
    "Not applicable. This manuscript does not reproduce copyrighted material from other "
    "sources."
)
ORIGINALITY = (
    "The author confirms that the work described has not been published previously, that "
    "it is not under consideration for publication elsewhere, and that, if accepted, it "
    "will not be published elsewhere in the same form, in English or in any other "
    "language, including electronically without the written consent of the copyright-holder."
)

REDACTIONS = [
    ("Altay's", "Author 1's"),
    ("Altay’s", "Author 1’s"),
    ("Yüksel, D., Altay, M.,", "Yüksel, D., Author 1,"),
    ("Altay, M.", "Author 1"),
    ("(Altay,", "(Author 1,"),
    (" Altay (", " Author 1 ("),
    (" Altay, ", " Author 1, "),
]


def configure_document(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def add_line(doc: Document, text: str = "", *, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(6)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_heading_line(doc: Document, text: str) -> None:
    add_line(doc, text, bold=True)


def build_title_page_document() -> Document:
    doc = Document()
    configure_document(doc)

    add_heading_line(doc, "Title")
    add_line(doc, TITLE)
    add_line(doc, "")

    add_heading_line(doc, "Author details")
    add_line(doc, "Corresponding author: Mehmet Altay")
    add_line(doc, "Position: Associate Professor")
    add_line(doc, "Affiliation: Kocaeli University, Türkiye")
    add_line(doc, "Department: Department of Foreign Language Education, Faculty of Education")
    add_line(doc, (
        "Address: Kabaoglu neighborhood, Baki Komsuoglu Boulevard, Kocaeli University, "
        "Umuttepe Campus, Faculty of Education, Department of Foreign Language Education, "
        "Postal code: 41100, Kocaeli, Türkiye"
    ))
    add_line(doc, "Email: mehmet.altay@kocaeli.edu.tr")
    add_line(doc, "ORCID ID: https://orcid.org/0000-0001-7227-5685")
    add_line(doc, "")

    add_line(doc, f"Keywords: {KEYWORDS}")
    add_line(doc, "")

    add_heading_line(doc, "Data availability statement")
    add_line(doc, DATA_AVAILABILITY)
    add_line(doc, "")

    add_heading_line(doc, "Funding statement")
    add_line(doc, FUNDING)
    add_line(doc, "")

    add_heading_line(doc, "Conflict of interest disclosure")
    add_line(doc, CONFLICT_OF_INTEREST)
    add_line(doc, "")

    add_heading_line(doc, "Ethics approval statement")
    add_line(doc, ETHICS_APPROVAL)
    add_line(doc, "")

    add_heading_line(doc, "Permission to reproduce material from other sources")
    add_line(doc, PERMISSION_TO_REPRODUCE)
    add_line(doc, "")

    add_heading_line(doc, "Originality and submission confirmation")
    add_line(doc, ORIGINALITY)
    add_line(doc, "")

    add_heading_line(doc, "CRediT author statement")
    add_line(doc, (
        "Mehmet Altay: Conceptualization, Methodology, Software, Validation, Formal Analysis, "
        "Investigation, Resources, Data Curation, Writing – Original Draft, Writing – Review and "
        "Editing, Visualization, Supervision, Project Administration."
    ))

    return doc


def save_title_page(path: Path) -> None:
    doc = build_title_page_document()
    doc.save(path)


def insert_title_page_at_start(source: Path, title_page: Document, destination: Path) -> None:
    doc = Document(str(source))
    body = doc.element.body

    prefix_elements = []
    for child in title_page.element.body:
        if child.tag.endswith("sectPr"):
            continue
        prefix_elements.append(deepcopy(child))

    page_break = doc.add_paragraph()
    page_break_run = page_break.add_run()
    page_break_run.add_break(WD_BREAK.PAGE)
    prefix_elements.append(deepcopy(page_break._element))

    for index, element in enumerate(prefix_elements):
        body.insert(index, element)

    doc.save(destination)


def build_manuscript_with_title_page(source: Path, destination: Path) -> None:
    title_page = build_title_page_document()
    insert_title_page_at_start(source, title_page, destination)


def redact_text(text: str) -> str:
    updated = text
    for old, new in REDACTIONS:
        updated = updated.replace(old, new)
    return re.sub(r"\bAltay\b", "Author 1", updated)


def redact_runs(container) -> None:
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            if run.text:
                run.text = redact_text(run.text)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                redact_runs(cell)


def build_blind_manuscript(source: Path, destination: Path) -> None:
    shutil.copy2(source, destination)
    doc = Document(str(destination))
    redact_runs(doc)
    for section in doc.sections:
        for header_footer in (section.header, section.footer):
            redact_runs(header_footer)
    doc.save(destination)


def verify_redaction(path: Path) -> None:
    doc = Document(str(path))
    text_parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.extend(paragraph.text for paragraph in cell.paragraphs)
    joined = "\n".join(text_parts)
    if re.search(r"\bAltay\b", joined):
        raise RuntimeError(f"Altay still present in {path.name}")
    if "Author 1" not in joined:
        raise RuntimeError(f"Expected Author 1 replacements missing in {path.name}")


def verify_merged_manuscript(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    required = [
        "Data availability statement",
        "Funding statement",
        "Conflict of interest disclosure",
        "Ethics approval statement",
        "Permission to reproduce material from other sources",
        "ORCID ID",
        TITLE,
        "Abstract",
    ]
    missing = [label for label in required if label not in text]
    if missing:
        raise RuntimeError(f"Merged manuscript missing: {', '.join(missing)}")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source manuscript not found: {SOURCE}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    title_page_path = OUTPUT_DIR / "Title page_With author details.docx"
    manuscript_path = OUTPUT_DIR / "IJAL v1.docx"
    blind_path = OUTPUT_DIR / "IJAL v1_Without author details.docx"

    save_title_page(title_page_path)
    build_manuscript_with_title_page(SOURCE, manuscript_path)
    build_blind_manuscript(SOURCE, blind_path)
    verify_redaction(blind_path)
    verify_merged_manuscript(manuscript_path)

    print(f"Created: {OUTPUT_DIR}")
    for item in sorted(OUTPUT_DIR.iterdir()):
        print(f"  - {item.name} ({item.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
