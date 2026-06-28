#!/usr/bin/env python3
"""Build TEAP final exam and answer key on Desktop."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DESKTOP = Path.home() / "Desktop"
EXAM_PATH = DESKTOP / "TEAP_Final_Test.docx"
KEY_PATH = DESKTOP / "TEAP_Final_Test_Answer_Key.docx"

FONT = "Times New Roman"
SIZE = Pt(11)


def set_run_font(run, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT
    run.font.size = SIZE
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def add_para(
    doc: Document,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    space_before: float = 0,
    space_after: float = 4,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    if text:
        run = p.add_run(text)
        set_run_font(run, bold=bold, italic=italic)


def add_mixed(doc: Document, parts: list[tuple[str, bool, bool]], space_after: float = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    for text, bold, italic in parts:
        run = p.add_run(text)
        set_run_font(run, bold=bold, italic=italic)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def add_header_block(doc: Document) -> None:
    add_para(doc, "Kocaeli University Faculty of Education TEFL Dept.", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "Final Test for TEAP Course", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_para(doc, "Date: 17/06/2006     Duration: 40 minutes     Total: 60 points", space_after=4)
    add_para(doc, "Student's Name / Surname: _________________________________________________", space_after=2)
    add_para(doc, "Student ID No.: ________________________________", space_after=8)


def build_exam() -> None:
    doc = Document()
    configure_doc(doc)
    add_header_block(doc)

    add_para(doc, "Course Participation Checklist", bold=True, space_before=4, space_after=4)
    add_para(doc, "Mark Yes or No for each item below.", space_after=6)

    checklist = [
        (
            "1.",
            "I participated in the Öğretim elemanlarının farklılıklara saygı düzeyleri questionnaire "
            "(Prof. Dr. Soner Polat).",
        ),
        ("2.", "I participated in the AI workshops (Büşra Kavan Alkan)."),
        ("3.", "I participated in the AI symposium at Kocaeli Sanayi Odası."),
    ]
    for num, text in checklist:
        add_para(doc, f"{num} {text}", space_after=2)
        add_para(doc, "     ☐ Yes     ☐ No", space_after=6)

    add_para(doc, "4. EAP presentation(s) in the classroom.", space_after=2)
    add_para(
        doc,
        "     ☐ I did not participate     ☐ I participated (number of times: __________)",
        space_after=6,
    )
    add_para(
        doc,
        "Instructor use only: participation verified  ☐ Yes     Signature: ____________________",
        space_after=10,
    )

    add_para(
        doc,
        "Instructions. Answer all questions. Write clearly. In Section I, write A for EAP or B for "
        "General English. In Section II, write the full name of the correct model. Section III "
        "requires short written answers.",
        space_after=8,
    )

    add_para(doc, "Section I. EAP vs. General English (10 points)", bold=True, space_after=4)
    add_para(
        doc,
        "For each statement, indicate whether it describes EAP (A) or General English (B).",
        space_after=6,
    )
    section_i = [
        "Course aims: Meet the needs of particular learners.",
        "Reason for study: Compulsory school course or adult evening class taken for interest.",
        "Course focus: Level-driven; begins with language.",
        "Timeframe: Short, fixed (e.g. a 4-12-week pre-sessional course).",
        "Stakes: High; university admission may depend on the result.",
        "Main skills focus: Listening and speaking receive more classroom time.",
        "Text type: Authentic, academic, genre-based texts explored in depth.",
        "Style of expression: Self expression and creativity are prized.",
        "Role of teacher: The teacher is the language expert.",
        "Learning skills: Strong emphasis on study skills, learner autonomy, and critical thinking.",
    ]
    for i, stem in enumerate(section_i, start=1):
        add_para(doc, f"{i}. {stem}     Answer: __________", space_after=3)

    add_para(doc, "Section II. Four Models of Putting Disciplines Together (12 points)", bold=True, space_before=8, space_after=4)
    add_para(
        doc,
        "Identify the model in each excerpt. Write Multidisciplinarity, Cross-disciplinarity, "
        "Interdisciplinarity, or Transdisciplinarity.",
        space_after=6,
    )
    section_ii = [
        (
            "11.",
            "At a research symposium on language education, linguistics, literature, translation studies, "
            "and education each hold separate sessions. Every session uses its own disciplinary questions "
            "and vocabulary; findings are presented side by side but not merged into a single new approach.",
        ),
        (
            "12.",
            "An EAP instructor applies a statistical technique developed in mathematics to interpret "
            "learners' placement-test scores. The instructor remains a language specialist; the borrowed "
            "tool simply helps answer a language-teaching question.",
        ),
        (
            "13.",
            "Psycholinguistics develops new constructs by blending methods, theories, and terminology from "
            "psychology and linguistics so that neither parent discipline alone could produce the resulting field.",
        ),
        (
            "14.",
            "To document an endangered language, university linguists work with community elders, local "
            "schoolteachers, and NGOs. Together they decide what to record, how materials will be used, and "
            "who owns the outcomes, because the real-world problem sets the research agenda.",
        ),
    ]
    for num, text in section_ii:
        add_para(doc, f"{num} {text}", space_after=2)
        add_para(doc, "Model: ________________________________", space_after=6)

    add_para(doc, "Section III. Transversal Skills (18 points)", bold=True, space_before=8, space_after=4)
    add_mixed(
        doc,
        [
            ("15. (9 points) ", True, False),
            ("According to UNESCO (UNESCO-UNEVOC / TVETipedia), how are ", False, False),
            ("transversal skills", False, True),
            (" defined? Write a concise definition in your own words (3-4 sentences).", False, False),
        ],
        space_after=6,
    )
    add_para(doc, "", space_after=14)
    add_para(doc, "", space_after=14)

    add_mixed(
        doc,
        [
            ("16. (9 points) ", True, False),
            ("UNESCO identifies six domains of transversal skills. Name all six domains and give ", False, False),
            ("one concrete example skill", True, False),
            (" for each domain (1.5 points per domain).", False, False),
        ],
        space_after=6,
    )
    for n in range(1, 7):
        add_para(doc, f"Domain {n}: ________________________________", space_after=2)
        add_para(doc, "Example: ________________________________", space_after=4)

    add_para(doc, "End of Test", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10)
    doc.save(EXAM_PATH)


def build_answer_key() -> None:
    doc = Document()
    configure_doc(doc)

    add_para(doc, "TEAP Course: Final Test Answer Key", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "Instructor copy only", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_para(doc, "Total: 60 points (checklist 20 + written test 40)", space_after=10)

    add_para(doc, "Participation Checklist (20 points)", bold=True, space_after=4)
    add_para(doc, "Each item below is worth 5 points unless noted.", space_after=6)

    checklist_key = [
        (
            "1. Polat questionnaire (5 pts)",
            "Yes = 5. No = 0. Accept only if the student completed the Öğretim elemanlarının "
            "farklılıklara saygı düzeyleri form.",
        ),
        (
            "2. AI workshops, Büşra Kavan Alkan (5 pts)",
            "Yes = 5. No = 0.",
        ),
        (
            "3. AI symposium, Kocaeli Sanayi Odası (5 pts)",
            "Yes = 5. No = 0.",
        ),
        (
            "4. EAP classroom presentation(s) (5 pts)",
            "Did not participate = 0. Participated once = 3. Participated twice or more = 5. "
            "Record the number the student writes.",
        ),
    ]
    for title, note in checklist_key:
        add_para(doc, title, bold=True, space_after=2)
        add_para(doc, note, space_after=6)

    add_para(doc, "Section I. EAP vs. General English (10 points; 1 point each)", bold=True, space_before=6, space_after=4)
    answers_i = [
        ("1.", "A (EAP)", "Meet the needs of particular learners."),
        ("2.", "B (General English)", "Compulsory course or interest."),
        ("3.", "B (General English)", "Level-driven; begins with language."),
        ("4.", "A (EAP)", "Short, fixed timeframe."),
        ("5.", "A (EAP)", "High stakes."),
        ("6.", "B (General English)", "More listening and speaking."),
        ("7.", "A (EAP)", "Authentic, academic, genre-based texts."),
        ("8.", "B (General English)", "Self expression and creativity."),
        ("9.", "B (General English)", "Teacher as language expert."),
        ("10.", "A (EAP)", "Study skills, autonomy, critical thinking."),
    ]
    for num, ans, note in answers_i:
        add_para(doc, f"{num} {ans}. {note}", space_after=2)

    add_para(doc, "Section II. Four Models (12 points; 3 points each)", bold=True, space_before=8, space_after=4)
    answers_ii = [
        ("11.", "Multidisciplinarity", "Separate sessions; juxtaposed, not synthesized."),
        ("12.", "Cross-disciplinarity", "One field borrows a tool from another."),
        ("13.", "Interdisciplinarity", "Blended methods/theories; new field emerges."),
        ("14.", "Transdisciplinarity", "Academics and non-academic actors co-produce knowledge."),
    ]
    for num, ans, note in answers_ii:
        add_para(doc, f"{num} {ans}. {note}", space_after=3)

    add_para(doc, "Section III. Transversal Skills (18 points)", bold=True, space_before=8, space_after=4)

    add_para(doc, "15. Definition (9 points)", bold=True, space_after=2)
    add_para(
        doc,
        "Award up to 9 points for a definition in the student's own words that includes the core UNESCO "
        "idea: skills not tied to one job, task, or discipline, and usable across situations and work "
        "settings. Partial credit (4-6) for incomplete but partly accurate answers. 0-3 if off-topic.",
        space_after=6,
    )
    add_para(
        doc,
        "Sample acceptable wording: Transversal skills are learned abilities that are not limited to a "
        "single subject area or occupation. They can be applied in many contexts, including study, "
        "work, and everyday life.",
        space_after=8,
    )

    add_para(doc, "16. Six domains with examples (9 points; 1.5 points per domain)", bold=True, space_after=2)
    domains = [
        ("1. Critical and innovative thinking", "e.g. reflective thinking, reasoned decision-making, creativity"),
        ("2. Interpersonal skills", "e.g. teamwork, communication, empathy"),
        ("3. Intrapersonal skills", "e.g. self-discipline, adaptability, perseverance"),
        ("4. Global citizenship", "e.g. tolerance, intercultural understanding, civic responsibility"),
        ("5. Media and information literacy", "e.g. locating information, critical evaluation of sources, ethical ICT use"),
        ("6. Physical and mental well-being", "e.g. healthy lifestyle, emotional self-regulation"),
    ]
    for domain, example in domains:
        add_para(doc, f"{domain}. Example: {example}", space_after=3)
    add_para(
        doc,
        "Scoring: 1 point for a correct domain name + 0.5 for a relevant example. Accept reasonable "
        "near-equivalents (e.g. 'collaboration' for teamwork).",
        space_after=8,
    )

    add_para(doc, "Score summary", bold=True, space_after=4)
    add_para(doc, "Checklist: ___ / 20", space_after=2)
    add_para(doc, "Section I: ___ / 10", space_after=2)
    add_para(doc, "Section II: ___ / 12", space_after=2)
    add_para(doc, "Section III: ___ / 18", space_after=2)
    add_para(doc, "Total: ___ / 60", bold=True, space_after=4)

    doc.save(KEY_PATH)


def main() -> None:
    build_exam()
    build_answer_key()
    print(f"Wrote {EXAM_PATH}")
    print(f"Wrote {KEY_PATH}")


if __name__ == "__main__":
    main()
