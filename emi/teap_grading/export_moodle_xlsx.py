#!/usr/bin/env python3
"""Merge TEAP grades into Moodle roster CSVs and save as Desktop xlsx."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from load_roster import ROSTER_CSV_ENCODING
from merge_results import EXAM_ABSENT, EXAM_RAW_MAX, scale_to_100

DOWNLOADS = Path.home() / "Downloads"
DESKTOP = Path.home() / "Desktop"
GRADES_CSV = DESKTOP / "TEAP_Final_Grades.csv"
OTHER_CSV = DESKTOP / "TEAP_Other_Courses_Exam.csv"

MOODLE_EXPORTS = [
    {
        "csv": DOWNLOADS / "1096320_Akademikİngilizce.csv",
        "xlsx": DESKTOP / "1096320_Akademikİngilizce.xlsx",
    },
    {
        "csv": DOWNLOADS / "1111747_Akademikİngilizce.csv",
        "xlsx": DESKTOP / "1111747_Akademikİngilizce.xlsx",
    },
]

GRADE_HEADERS = [
    ("participation", "Katılım (20)"),
    ("section_i", "Bölüm I (10)"),
    ("section_ii", "Bölüm II (12)"),
    ("section_iii", "Bölüm III (18)"),
    ("total", "Toplam (60)"),
    ("total_100", "Toplam (100)"),
    ("exam_status", "Sınav"),
]


def load_grades() -> dict[str, dict]:
    grades: dict[str, dict] = {}
    with GRADES_CSV.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            sid = row["idnumber"].strip()
            absent = row.get("total") == EXAM_ABSENT
            if absent:
                grades[sid] = {
                    "participation": EXAM_ABSENT,
                    "section_i": EXAM_ABSENT,
                    "section_ii": EXAM_ABSENT,
                    "section_iii": EXAM_ABSENT,
                    "total": EXAM_ABSENT,
                    "total_100": EXAM_ABSENT,
                    "exam_status": EXAM_ABSENT,
                    "notes": row.get("notes", ""),
                }
            else:
                total = int(float(row["total"]))
                grades[sid] = {
                    "participation": int(float(row["participation"])),
                    "section_i": int(float(row["section_i"])),
                    "section_ii": int(float(row["section_ii"])),
                    "section_iii": int(float(row["section_iii"])),
                    "total": total,
                    "total_100": scale_to_100(total),
                    "exam_status": "",
                    "notes": row.get("notes", ""),
                }
    return grades


def autosize_columns(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 48)


def write_xlsx(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Notlar"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    autosize_columns(ws)
    wb.save(path)


def export_moodle_file(csv_path: Path, xlsx_path: Path, grades: dict[str, dict]) -> int:
    with csv_path.open(encoding=ROSTER_CSV_ENCODING) as f:
        reader = csv.DictReader(f, delimiter=";")
        base_fields = list(reader.fieldnames or [])
        moodle_rows = list(reader)

    extra_headers = [label for _, label in GRADE_HEADERS] + ["Not açıklaması"]
    headers = base_fields + extra_headers
    out_rows: list[list[object]] = []
    matched = 0

    for row in moodle_rows:
        sid = row["idnumber"].strip()
        g = grades.get(sid)
        if g:
            matched += 1
        else:
            g = {
                "participation": "",
                "section_i": "",
                "section_ii": "",
                "section_iii": "",
                "total": "",
                "total_100": "",
                "exam_status": "",
                "notes": "Not bulunamadı",
            }
        out_rows.append(
            [row.get(field, "") for field in base_fields]
            + [g.get(key, "") for key, _ in GRADE_HEADERS]
            + [g.get("notes", "")]
        )

    write_xlsx(xlsx_path, headers, out_rows)
    return matched


def export_other_courses() -> None:
    if not OTHER_CSV.exists():
        return
    headers = [
        "idnumber",
        "Ad (kağıt)",
        "Grup",
        "Katılım (20)",
        "Bölüm I (10)",
        "Bölüm II (12)",
        "Bölüm III (18)",
        "Toplam (60)",
        "Katılım (100)",
        "Bölüm I (100)",
        "Bölüm II (100)",
        "Bölüm III (100)",
        "Toplam (100)",
        "Klasör",
        "Not açıklaması",
    ]
    rows: list[list[object]] = []
    with OTHER_CSV.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            rows.append(
                [
                    row["idnumber"],
                    row.get("name_on_paper", ""),
                    row.get("course_group", "diğer dersler"),
                    row["participation"],
                    row["section_i"],
                    row["section_ii"],
                    row["section_iii"],
                    row["total"],
                    row.get("participation_100", ""),
                    row.get("section_i_100", ""),
                    row.get("section_ii_100", ""),
                    row.get("section_iii_100", ""),
                    row.get("total_100", ""),
                    row.get("folder", ""),
                    row.get("notes", ""),
                ]
            )
    write_xlsx(DESKTOP / "TEAP_Diger_Dersler.xlsx", headers, rows)


def main() -> None:
    if not GRADES_CSV.exists():
        raise SystemExit(f"Missing grades file: {GRADES_CSV}")

    grades = load_grades()
    for item in MOODLE_EXPORTS:
        csv_path = item["csv"]
        xlsx_path = item["xlsx"]
        if not csv_path.exists():
            raise SystemExit(f"Missing roster CSV: {csv_path}")
        n = export_moodle_file(csv_path, xlsx_path, grades)
        print(f"Wrote {xlsx_path} ({n} students with grades)")

    export_other_courses()
    if OTHER_CSV.exists():
        print(f"Wrote {DESKTOP / 'TEAP_Diger_Dersler.xlsx'}")

    print(f"Scale: {EXAM_RAW_MAX} → 100 orantılı, tam sayı")


if __name__ == "__main__":
    main()
