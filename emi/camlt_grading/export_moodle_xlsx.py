#!/usr/bin/env python3
"""Export CAMLT grades to Moodle xlsx.

Formats:
- full (default): İsim, Soyisim, Öğrenci numarası, Total puan — with header row
- moodle-upload: Öğrenci numarası, final puan only — no header row (Moodle bulk import)

Upload files on Desktop:
  CAM 1 - Moodle yüklemesi.xlsx  (group 1096037)
  CAM 2 - Moodle yüklemesi.xlsx  (group 1097165)
  CAM 3 - Moodle yüklemesi.xlsx  (group 1105959)
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from load_roster import ROSTER_CSV_ENCODING, fix_turkish_mojibake

DOWNLOADS = Path.home() / "Downloads"
DESKTOP = Path.home() / "Desktop"
GRADES_CSV = DESKTOP / "CAMLT_Final_Grades.csv"
EXAM_ABSENT = "E"

# Confirmed absent — always keep E regardless of grade file content.
CONFIRMED_ABSENT: frozenset[str] = frozenset(
    {
        "230907033",  # Sude Öksüz
        "220907004",  # Zeliha Kars
        "250921047",  # Muhammet Taha Kisbet
        "210907031",  # Doğukan Aysan
        "200907052",  # İbrahim Ozan Özmener
    }
)

MOODLE_EXPORTS = [
    {
        "csv": DOWNLOADS / "1096037_DilÖğrenmeveÖğretimYaklaşımlarıII.csv",
        "xlsx": DESKTOP / "CAM 1.xlsx",
        "upload_xlsx": DESKTOP / "CAM 1 - Moodle yüklemesi.xlsx",
        "group": "1096037",
        "label": "CAM 1",
    },
    {
        "csv": DOWNLOADS / "1097165_DilÖğrenmeveÖğretimYaklaşımlarıII.csv",
        "xlsx": DESKTOP / "CAM 2.xlsx",
        "upload_xlsx": DESKTOP / "CAM 2 - Moodle yüklemesi.xlsx",
        "group": "1097165",
        "label": "CAM 2",
    },
    {
        "csv": DOWNLOADS / "1105959_İngilizceÖğretimProgramları.csv",
        "xlsx": DESKTOP / "CAM 3.xlsx",
        "upload_xlsx": DESKTOP / "CAM 3 - Moodle yüklemesi.xlsx",
        "group": "1105959",
        "label": "CAM 3",
    },
]

HEADERS = ["İsim", "Soyisim", "Öğrenci numarası", "Total puan"]


@dataclass
class ExportStats:
    group: str
    xlsx_path: Path
    roster_count: int
    graded_count: int
    absent_count: int
    missing_in_grades: list[str]


def load_grades() -> dict[str, str | int]:
    grades: dict[str, str | int] = {}
    with GRADES_CSV.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            sid = row["idnumber"].strip()
            if sid in CONFIRMED_ABSENT:
                grades[sid] = EXAM_ABSENT
                continue

            total_raw = (row.get("total") or "").strip()
            if total_raw == EXAM_ABSENT:
                grades[sid] = EXAM_ABSENT
            else:
                grades[sid] = int(round(float(total_raw)))
    return grades


def autosize_columns(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 48)


def write_xlsx(
    path: Path, rows: list[list[object]], *, headers: list[str] | None = None
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Notlar"
    if headers:
        ws.append(headers)
    for row in rows:
        ws.append(row)
    autosize_columns(ws)
    wb.save(path)


def resolve_total(grades: dict[str, str | int], sid: str) -> tuple[str | int, bool, bool]:
    """Return (total_value, is_graded, is_absent). Missing grades → empty string."""
    total = grades.get(sid)
    if total is None:
        return "", False, False
    if total == EXAM_ABSENT:
        return EXAM_ABSENT, False, True
    return total, True, False


def export_moodle_file(
    csv_path: Path,
    xlsx_path: Path,
    group: str,
    grades: dict[str, str | int],
    *,
    moodle_upload: bool = False,
) -> ExportStats:
    with csv_path.open(encoding=ROSTER_CSV_ENCODING) as f:
        reader = csv.DictReader(f, delimiter=";")
        moodle_rows = list(reader)

    out_rows: list[list[object]] = []
    graded_count = 0
    absent_count = 0
    missing_in_grades: list[str] = []

    for row in moodle_rows:
        sid = row["idnumber"].strip()
        total_value, is_graded, is_absent = resolve_total(grades, sid)

        if total_value == "":
            missing_in_grades.append(sid)
        elif is_absent:
            absent_count += 1
        elif is_graded:
            graded_count += 1

        if moodle_upload:
            out_rows.append([sid, total_value])
        else:
            first = fix_turkish_mojibake(row["firstname"].strip())
            last = fix_turkish_mojibake(row["lastname"].strip())
            out_rows.append([first, last, sid, total_value])

    write_xlsx(
        xlsx_path,
        out_rows,
        headers=None if moodle_upload else HEADERS,
    )
    return ExportStats(
        group=group,
        xlsx_path=xlsx_path,
        roster_count=len(moodle_rows),
        graded_count=graded_count,
        absent_count=absent_count,
        missing_in_grades=missing_in_grades,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export CAMLT grades to Moodle xlsx files on Desktop."
    )
    parser.add_argument(
        "--format",
        choices=("full", "moodle-upload", "both"),
        default="full",
        help=(
            "full: 4-column with headers (CAM 1.xlsx …); "
            "moodle-upload: öğrenci no + puan, başlıksız; "
            "both: write all six files"
        ),
    )
    return parser.parse_args()


def main() -> list[ExportStats]:
    args = parse_args()
    if not GRADES_CSV.exists():
        raise SystemExit(f"Missing grades file: {GRADES_CSV}")

    grades = load_grades()
    all_stats: list[ExportStats] = []
    write_full = args.format in ("full", "both")
    write_upload = args.format in ("moodle-upload", "both")

    for item in MOODLE_EXPORTS:
        csv_path = item["csv"]
        if not csv_path.exists():
            raise SystemExit(f"Missing roster CSV: {csv_path}")

        if write_full:
            stats = export_moodle_file(
                csv_path, item["xlsx"], item["group"], grades, moodle_upload=False
            )
            all_stats.append(stats)
            print(
                f"Wrote {stats.xlsx_path} "
                f"(roster={stats.roster_count}, graded={stats.graded_count}, E={stats.absent_count})"
            )
            if stats.missing_in_grades:
                print(f"  WARNING: not in grade file: {', '.join(stats.missing_in_grades)}")

        if write_upload:
            stats = export_moodle_file(
                csv_path,
                item["upload_xlsx"],
                item["group"],
                grades,
                moodle_upload=True,
            )
            all_stats.append(stats)
            print(
                f"Wrote {stats.xlsx_path} [başlıksız yükleme] "
                f"(roster={stats.roster_count}, graded={stats.graded_count}, E={stats.absent_count})"
            )
            if stats.missing_in_grades:
                print(f"  WARNING: not in grade file: {', '.join(stats.missing_in_grades)}")

    print(f"Confirmed absent (E): {', '.join(sorted(CONFIRMED_ABSENT))}")
    return all_stats


if __name__ == "__main__":
    main()
