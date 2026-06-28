#!/usr/bin/env python3
"""Import manual edits from CAMLT_Final_Grades_Detail.xlsx (highlights + total overrides)."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent
DEFAULT_XLSX = Path.home() / "Desktop" / "CAMLT_Final_Grades_Detail.xlsx"
OUT_CSV = Path.home() / "Desktop" / "CAMLT_Final_Grades_Detail.csv"
OUT_MAIN = Path.home() / "Desktop" / "CAMLT_Final_Grades.csv"
RESULTS_DIR = ROOT / "results"
REPORT = ROOT / "import_report.txt"

# Yellow / highlight-like fills (Excel theme + RGB).
HIGHLIGHT_RGB = {
    "FFFFFF00",
    "FFFF00",
    "FFFFFF99",
    "FFFFCC00",
    "FFFFEB9C",
    "FFFFFFCC",
    "FFF2CC",
    "FFFFCC",
    "FFFFFF66",
}


def is_highlighted(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False
    fg = fill.fgColor
    if fg is None:
        return False
    if fg.type == "rgb" and fg.rgb:
        rgb = fg.rgb[-6:].upper() if len(fg.rgb) >= 6 else fg.rgb.upper()
        if rgb in HIGHLIGHT_RGB or rgb.startswith("FFFF"):
            return True
    if fg.type == "indexed" and fg.indexed in (6, 13, 19, 43, 44):
        return True
    if fg.type == "theme" and fg.theme in (3, 4, 5, 6, 7):
        return True
    return False


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val is None:
            continue
        mapping[str(val).strip().lower()] = col
    return mapping


def load_csv_rows(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def save_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow(row)


def recompute_sections(row: dict) -> None:
    def num(k: str) -> float:
        v = row.get(k, "")
        if v in ("", "E", None):
            return 0.0
        return float(v)

    if row.get("total") == "E":
        return
    row["section1"] = num("q1a") + num("q1b") + num("q1c") + num("q1d")
    row["section2"] = num("q2")
    row["section3"] = num("q3a") + num("q3b") + num("q3c")


def update_batch_json(folder: str, updates: dict[str, float | int]) -> bool:
    if not folder:
        return False
    for path in sorted(RESULTS_DIR.glob("batch*.json")):
        rows = json.loads(path.read_text(encoding="utf-8"))
        changed = False
        for row in rows:
            if row.get("folder") != folder:
                continue
            for k, v in updates.items():
                if k in row and row[k] != v:
                    row[k] = v
                    changed = True
            if changed:
                row["section1"] = sum(row[x] for x in ("q1a", "q1b", "q1c", "q1d"))
                row["section2"] = row["q2"]
                row["section3"] = sum(row[x] for x in ("q3a", "q3b", "q3c"))
                row["total"] = int(
                    round(row["section1"] + row["section2"] + row["section3"])
                )
            path.write_text(
                json.dumps(rows, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            return changed
    return False


def main() -> None:
    xlsx_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}")
        print("Save your highlighted Excel as Desktop/CAMLT_Final_Grades_Detail.xlsx")
        sys.exit(1)

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    cols = header_map(ws)
    required = ["idnumber", "total"]
    for r in required:
        if r not in cols:
            print(f"Missing column '{r}' in row 1. Found: {list(cols)}")
            sys.exit(1)

    score_cols = [
        c
        for c in ("total", "section1", "section2", "section3", "q1a", "q1b", "q1c", "q1d", "q2", "q3a", "q3b", "q3c")
        if c in cols
    ]

    highlighted: list[str] = []
    total_changes: list[str] = []
    field_changes: list[str] = []

    csv_rows = load_csv_rows(OUT_CSV) if OUT_CSV.exists() else load_csv_rows(OUT_MAIN)
    by_id = {r["idnumber"]: r for r in csv_rows if r.get("idnumber")}

    for row_idx in range(2, ws.max_row + 1):
        sid_cell = ws.cell(row_idx, cols["idnumber"])
        sid = sid_cell.value
        if sid is None or str(sid).strip() == "":
            continue
        sid = str(sid).strip()
        row = by_id.get(sid)
        if not row:
            continue

        folder = row.get("folder", "")
        json_updates: dict[str, float | int] = {}
        row_notes: list[str] = []

        for col_name in score_cols:
            col_i = cols[col_name]
            cell = ws.cell(row_idx, col_i)
            val = cell.value
            if val is None or val == "":
                continue
            if col_name == "total" and str(val).upper() == "E":
                new_val = "E"
            else:
                try:
                    new_val = int(round(float(val)))
                except (TypeError, ValueError):
                    continue

            old_raw = row.get(col_name, "")
            old = old_raw if old_raw == "E" else int(round(float(old_raw or 0)))
            if new_val != old:
                row[col_name] = new_val
                if col_name == "total":
                    total_changes.append(f"{sid} {row.get('roster_fullname','')}: {old} -> {new_val}")
                else:
                    field_changes.append(f"{sid} {col_name}: {old} -> {new_val}")
                    if col_name in ("q1a", "q1b", "q1c", "q1d", "q2", "q3a", "q3b", "q3c"):
                        json_updates[col_name] = float(new_val)

            if is_highlighted(cell):
                row_notes.append(col_name)

        if row_notes:
            highlighted.append(
                f"Row {row_idx} {sid} {row.get('roster_fullname','')}: highlighted {', '.join(row_notes)}"
            )

        if json_updates and folder:
            update_batch_json(folder, json_updates)
        if row.get("total") != "E":
            recompute_sections(row)

    fieldnames = list(csv_rows[0].keys()) if csv_rows else []
    roster_only = [r for r in csv_rows if r.get("match_status") != "unmatched_paper" and r.get("idnumber")]
    save_csv(OUT_CSV, csv_rows, fieldnames)
    save_csv(OUT_MAIN, roster_only, fieldnames)

    lines = [
        f"Source: {xlsx_path}",
        f"Highlighted cells: {len(highlighted)}",
        *highlighted,
        "",
        f"Total changes: {len(total_changes)}",
        *total_changes,
        "",
        f"Other field changes: {len(field_changes)}",
        *field_changes,
    ]
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines[:30]))
    if len(lines) > 30:
        print(f"... full report: {REPORT}")
    print(f"\nWrote {OUT_CSV} and {OUT_MAIN}")


if __name__ == "__main__":
    main()
