#!/usr/bin/env python3
"""Reform the quantitative workbook to pre/post proficiency comparison only.

Keeps paired skill scores + light instrument reliability.
Drops Golem / mechanism / regression / mediation sheets from the primary workbook.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from scipy import stats

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "EMI_PYP_pre_post_synthetic_N120.xlsx"
OUT = ROOT / "EMI_quantitative_prepost_N120.xlsx"
DESK = Path("/home/ubuntu/Desktop/EMI Discouragement Project")

CORE_COLS = [
    "Student_ID",
    "Major",
    "Gender",
    "Age_at_pre_PYP_exit",
    "Age_at_post_graduation",
    "Pre_Listening",
    "Pre_Reading",
    "Pre_Writing",
    "Pre_Speaking",
    "Pre_Overall",
    "Post_Listening",
    "Post_Reading",
    "Post_Writing",
    "Post_Speaking",
    "Post_Overall",
    "Decline_Listening",
    "Decline_Reading",
    "Decline_Writing",
    "Decline_Speaking",
    "Decline_Overall",
]

SKILLS = ["Listening", "Reading", "Writing", "Speaking", "Overall"]


def paired_table(s: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for sk in SKILLS:
        pre, post = s[f"Pre_{sk}"], s[f"Post_{sk}"]
        d = pre - post
        t = stats.ttest_rel(pre, post)
        w = stats.wilcoxon(pre, post, zero_method="wilcox", alternative="two-sided")
        se = d.std(ddof=1) / np.sqrt(len(d))
        ci = stats.t.interval(0.95, len(d) - 1, loc=d.mean(), scale=se)
        dz = d.mean() / d.std(ddof=1)
        sw = stats.shapiro(d)
        rows.append(
            {
                "Skill": sk,
                "N": len(s),
                "Pre_M": pre.mean(),
                "Pre_SD": pre.std(ddof=1),
                "Post_M": post.mean(),
                "Post_SD": post.std(ddof=1),
                "Mdiff_pre_minus_post": d.mean(),
                "SD_diff": d.std(ddof=1),
                "t": t.statistic,
                "df": len(s) - 1,
                "p_t": t.pvalue,
                "d_z": dz,
                "CI95_lo": ci[0],
                "CI95_hi": ci[1],
                "Wilcoxon_W": w.statistic,
                "p_Wilcoxon": w.pvalue,
                "Shapiro_W": sw.statistic,
                "Shapiro_p": sw.pvalue,
                "Sig_decline": "yes" if (d.mean() > 0 and t.pvalue < 0.05) else "no",
            }
        )
    return pd.DataFrame(rows)


def descriptives(s: pd.DataFrame) -> pd.DataFrame:
    cols = [c for c in CORE_COLS if c.startswith(("Pre_", "Post_", "Decline_"))]
    rows = []
    for c in cols:
        x = s[c]
        rows.append(
            {
                "Variable": c,
                "N": len(x),
                "M": x.mean(),
                "SD": x.std(ddof=1),
                "Min": x.min(),
                "Max": x.max(),
                "Skew": float(x.skew()),
            }
        )
    return pd.DataFrame(rows)


def write_readme(ws) -> None:
    lines = [
        "EMI quantitative panel — Pre/Post proficiency only (synthetic N = 120)",
        "",
        "Scope: RQ1 only — paired comparison of institutional English scores",
        "(Listening, Reading, Writing, Speaking, Overall) from PYP exit to graduation.",
        "Explanatory accounts of attrition (translanguaging, lecturer/student perspectives)",
        "belong to the QUALITATIVE strand (RQ2–RQ3) and are not in this workbook.",
        "",
        "Scale: 0–100 institutional; PYP pass threshold = 60 (≈ CEFR B1).",
        "Decline = Pre − Post (positive = attrition).",
        "",
        "Sheets: README, Codebook, Students, Descriptives, Paired_pre_post,",
        "Writing_Pre_raters, Writing_Post_raters, Speaking_Pre_raters, Speaking_Post_raters, ICC_raters.",
        "",
        "Replace synthetic values with live institutional data before submission.",
    ]
    ws["A1"] = "README"
    ws["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(lines, start=3):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 100


def write_codebook(ws) -> None:
    rows = [
        ("Student_ID", "Participant id"),
        ("Major", "Engineering major"),
        ("Gender", "Gender"),
        ("Age_at_pre_PYP_exit", "Age at PYP-exit test"),
        ("Age_at_post_graduation", "Age at graduation test"),
        ("Pre_*", "PYP-exit skill / overall (0–100)"),
        ("Post_*", "Graduation skill / overall (0–100)"),
        ("Decline_*", "Pre − Post (positive = attrition)"),
        ("PYP_threshold", "60/100 ≈ CEFR B1"),
    ]
    ws["A1"] = "Variable"
    ws["B1"] = "Definition"
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    for i, (a, b) in enumerate(rows, start=2):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 55


def df_to_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(name)
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(1, j, col)
        cell.font = Font(bold=True)
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, val in enumerate(row, start=1):
            if isinstance(val, float):
                ws.cell(i, j, round(val, 4) if abs(val) < 1000 else val)
            else:
                ws.cell(i, j, val)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(18, 4 + max(len(str(c.value or "")) for c in col))


def main() -> None:
    students_full = pd.read_excel(SRC, "Students")
    students = students_full[CORE_COLS].copy()

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    write_readme(ws0)
    write_codebook(wb.create_sheet("Codebook"))
    df_to_sheet(wb, "Students", students)
    df_to_sheet(wb, "Descriptives", descriptives(students))
    df_to_sheet(wb, "Paired_pre_post", paired_table(students))

    # Keep rater sheets if present (instrument quality only)
    for sheet in [
        "Writing_Pre_raters",
        "Writing_Post_raters",
        "Speaking_Pre_raters",
        "Speaking_Post_raters",
        "ICC_raters",
    ]:
        try:
            df = pd.read_excel(SRC, sheet)
            df_to_sheet(wb, sheet, df)
        except Exception:
            pass

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    # Also overwrite the legacy filename used by figure/doc scripts for compatibility
    legacy = ROOT / "EMI_PYP_pre_post_synthetic_N120.xlsx"
    # Keep a mechanism archive copy once, then replace primary with slim? 
    # Safer: write slim OUT and also save slim as primary quant path used by docs.
    DESK.mkdir(parents=True, exist_ok=True)
    wb.save(DESK / OUT.name)
    # Point legacy name to slim quant for figure scripts that still load SRC name
    wb.save(legacy)
    wb.save(DESK / legacy.name)

    print("wrote", OUT)
    print(paired_table(students)[["Skill", "Mdiff_pre_minus_post", "t", "p_t", "d_z", "Sig_decline"]].to_string(index=False))


if __name__ == "__main__":
    main()
